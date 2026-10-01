from __future__ import annotations

import io
import logging
import os

logger = logging.getLogger(__name__)

MAX_TEXTO_ARQUIVO = 5000

_TESSERACT_OK: bool | None = None
_OCR_LANG: str | None = None

_TESSERACT_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
]


def _configurar_tesseract() -> None:
    import pytesseract

    for caminho in _TESSERACT_PATHS:
        if os.path.exists(caminho):
            pytesseract.pytesseract.tesseract_cmd = caminho
            return


def _tesseract_disponivel() -> bool:
    global _TESSERACT_OK
    if _TESSERACT_OK is None:
        try:
            import pytesseract
            _configurar_tesseract()
            pytesseract.get_tesseract_version()
            _TESSERACT_OK = True
            logger.info("Tesseract disponivel — OCR habilitado")
        except Exception:
            _TESSERACT_OK = False
            logger.warning("Tesseract nao disponivel — OCR desabilitado")
    return _TESSERACT_OK


def _ocr_lang() -> str:
    global _OCR_LANG
    if _OCR_LANG is None:
        try:
            import pytesseract
            disponiveis = set(pytesseract.get_languages() or [])
            escolhidos = [l for l in ("por", "eng") if l in disponiveis]
            _OCR_LANG = "+".join(escolhidos)
        except Exception:
            _OCR_LANG = ""
    return _OCR_LANG


def _ocr_imagem(data: bytes) -> str:
    import pytesseract
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    return pytesseract.image_to_string(img, lang=_ocr_lang() or None)


def extrair_texto(ext: str, data: bytes) -> str | None:
    ext = (ext or "").lower().lstrip(".")
    texto = None
    try:
        if ext == "txt":
            texto = _ler_texto(data)
        elif ext == "pdf":
            texto = _extrair_pdf(data)
            if not (texto and texto.strip()) and _tesseract_disponivel():
                texto = _ocr_pdf(data)
        elif ext == "docx":
            texto = _extrair_docx(data)
        elif ext == "pptx":
            texto = _extrair_pptx(data)
        elif ext in ("png", "jpg", "jpeg"):
            if _tesseract_disponivel():
                texto = _ocr_imagem(data)
    except Exception as exc:
        logger.warning("Falha ao extrair texto de .%s: %s", ext, exc)
        return None

    texto = (texto or "").strip()
    if not texto:
        return None

    if len(texto) > MAX_TEXTO_ARQUIVO:
        texto = texto[:MAX_TEXTO_ARQUIVO] + "\n...[texto truncado]"

    return texto


def _ler_texto(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


def _extrair_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    partes = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            partes.append(t.strip())
    return "\n\n".join(partes)


def _ocr_pdf(data: bytes) -> str:
    partes = []
    try:
        import fitz
    except ImportError:
        fitz = None

    if fitz is not None:
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            zoom = 200 / 72
            mat = fitz.Matrix(zoom, zoom)
            for page in doc:
                pix = page.get_pixmap(matrix=mat)
                texto = _ocr_imagem(pix.tobytes("png"))
                if texto and texto.strip():
                    partes.append(texto.strip())
            doc.close()
            if partes:
                return "\n\n".join(partes)
        except Exception as exc:
            logger.warning("Falha no OCR via pymupdf: %s", exc)

    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    for page in reader.pages:
        try:
            for img in page.images:
                try:
                    texto = _ocr_imagem(img.data)
                except Exception as exc:
                    logger.warning("Falha no OCR de imagem do PDF: %s", exc)
                    continue
                if texto and texto.strip():
                    partes.append(texto.strip())
        except Exception as exc:
            logger.warning("Falha ao ler imagens da pagina: %s", exc)
            continue
    return "\n\n".join(partes)


def _extrair_docx(data: bytes) -> str:
    import docx

    documento = docx.Document(io.BytesIO(data))
    paragrafos = [p.text for p in documento.paragraphs if p.text and p.text.strip()]
    return "\n".join(paragrafos)


def _extrair_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    partes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text:
                partes.append(shape.text.strip())
    return "\n".join(partes)